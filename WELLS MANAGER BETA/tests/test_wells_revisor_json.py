# -*- coding: utf-8 -*-
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
JSON_DIR = ROOT / "Wells_Revisor_JSON"
if str(JSON_DIR) not in sys.path:
    sys.path.insert(0, str(JSON_DIR))

from wells_blindagem_json import (
    LISTA_BLINDADOS, LISTA_TRADUZIVEIS, salvar_listas
)
from wells_revisor_json_core import (
    review_json_pair, diagnosticar_par_linguistico,
    _correct_source_anchored_structures,
)
from wells_isolador_renpy_json import segmentar


class RevisorJsonIndependenteTests(unittest.TestCase):
    def _pair(self, native, translated, blindados=(), traduziveis=()):
        td = tempfile.TemporaryDirectory()
        root = Path(td.name)
        n = root / "nativo.txt"
        t = root / "traduzido.txt"
        b = root / "blindagem.txt"
        n.write_text("\n".join(f'"{x}"' for x in native) + "\n", encoding="utf-8")
        t.write_text("\n".join(f'"{x}"' for x in translated) + "\n", encoding="utf-8")
        salvar_listas(b, {
            LISTA_BLINDADOS: list(blindados),
            LISTA_TRADUZIVEIS: list(traduziveis),
        })
        before = n.read_bytes()
        result = review_json_pair(n, t, root, b)
        after = n.read_bytes()
        lines = t.read_text(encoding="utf-8").splitlines()
        return td, result, before, after, lines

    def test_blindado_exato_e_restaurado_quando_sobrevive(self):
        td, result, before, after, lines = self._pair(
            ["Christine"], ["christine"], blindados=["Christine"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Christine"'])
        self.assertEqual(result["protected_names"], 1)
        self.assertEqual(result["unresolved_names"], [])
        self.assertEqual(before, after)

    def test_blindagem_age_dentro_de_registro_composto_sem_congelar_unidade(self):
        td, result, before, after, lines = self._pair(
            ["Princess Christine"], ["Princesa christine"],
            blindados=["Christine"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Princesa Christine"'])
        self.assertEqual(result["protected_records"], 1)
        self.assertEqual(result["protected_names"], 1)
        self.assertEqual(before, after)

    def test_termo_maior_nao_casa_com_parte_de_outra_palavra(self):
        td, result, before, after, lines = self._pair(
            ["Princess Dalenor"], ["Princesa dalenor"],
            blindados=["Dalen"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Princesa dalenor"'])
        self.assertEqual(result["protected_names"], 0)
        self.assertEqual(result["unresolved_names"], [])
        self.assertEqual(before, after)

    def test_nome_destruido_em_registro_integral_e_restaurado(self):
        td, result, before, after, lines = self._pair(
            ["Dalen"], ["O Vale"], blindados=["Dalen"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Dalen"'])
        self.assertEqual(result["protected_names"], 1)
        self.assertEqual(result["reconstructed_names"], 1)
        self.assertEqual(result["unresolved_names"], [])
        self.assertEqual(before, after)

    def test_nome_destruido_em_composto_sem_classificacao_total_fica_pendente(self):
        td, result, before, after, lines = self._pair(
            ["Commander Gaius arrives"], ["Comandante Go chega"],
            blindados=["Gaius"], traduziveis=["Commander"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Comandante Go chega"'])
        self.assertEqual(result["protected_names"], 0)
        self.assertEqual(result["reconstructed_names"], 0)
        self.assertEqual(len(result["unresolved_names"]), 1)
        self.assertEqual(result["unresolved_names"][0]["term"], "Gaius")
        self.assertEqual(before, after)

    def test_nome_composto_reconstroi_com_slots_manuais_deterministicos(self):
        td, result, before, after, lines = self._pair(
            ["Commander Gaius"], ["Comandante Go"],
            blindados=["Gaius"], traduziveis=["Commander"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Comandante Gaius"'])
        self.assertEqual(result["protected_names"], 1)
        self.assertEqual(result["reconstructed_names"], 1)
        self.assertEqual(result["unresolved_names"], [])
        self.assertEqual(before, after)

    def test_composto_sem_ancora_traduzivel_continua_pendente(self):
        td, result, before, after, lines = self._pair(
            ["Commander Gaius"], ["Comandante Go"],
            blindados=["Gaius"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Comandante Go"'])
        self.assertEqual(result["reconstructed_names"], 0)
        self.assertEqual(len(result["unresolved_names"]), 1)
        self.assertEqual(before, after)

    def test_composto_com_contagem_lexical_diferente_continua_pendente(self):
        td, result, before, after, lines = self._pair(
            ["Commander Gaius"], ["O grande comandante Go"],
            blindados=["Gaius"], traduziveis=["Commander"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"O grande comandante Go"'])
        self.assertEqual(result["reconstructed_names"], 0)
        self.assertEqual(len(result["unresolved_names"]), 1)
        self.assertEqual(before, after)

    def test_composto_com_nativo_nao_totalmente_classificado_continua_pendente(self):
        td, result, before, after, lines = self._pair(
            ["Supreme Commander Gaius"], ["Supremo Comandante Go"],
            blindados=["Gaius"], traduziveis=["Commander"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Supremo Comandante Go"'])
        self.assertEqual(result["reconstructed_names"], 0)
        self.assertEqual(len(result["unresolved_names"]), 1)
        self.assertEqual(before, after)

    def test_lista_vazia_nao_cria_protecao_implicita(self):
        td, result, before, after, lines = self._pair(
            ["Dalen"], ["O Vale"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"O Vale"'])
        self.assertEqual(result["protected_names"], 0)
        self.assertEqual(result["unresolved_names"], [])
        self.assertEqual(before, after)

    def test_multiplos_nomes_blindados_sobreviventes_sao_preservados(self):
        td, result, before, after, lines = self._pair(
            ["Dalen meets Gaius"], ["dalen encontra gaius"],
            blindados=["Dalen", "Gaius"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Dalen encontra Gaius"'])
        self.assertEqual(result["protected_names"], 2)
        self.assertEqual(result["unresolved_names"], [])
        self.assertEqual(before, after)

    def test_overlap_ann_nao_altera_anna(self):
        td, result, before, after, lines = self._pair(
            ["Anna meets Ann"], ["anna encontra ann"],
            blindados=["Ann", "Anna"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Anna encontra Ann"'])
        self.assertEqual(result["protected_names"], 2)
        self.assertEqual(before, after)

    def test_nome_acentuado_reconstruido_por_slot_classificado(self):
        td, result, before, after, lines = self._pair(
            ["Commander Éloïse"], ["Comandante Eloise"],
            blindados=["Éloïse"], traduziveis=["Commander"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Comandante Éloïse"'])
        self.assertEqual(result["reconstructed_names"], 1)
        self.assertEqual(result["unresolved_names"], [])
        self.assertEqual(before, after)

    def test_apostrofo_em_nome_sobrevivente_e_reconstruido(self):
        td, result, before, after, lines = self._pair(
            ["Commander A'nai", "Commander A'nai"],
            ["Comandante a'nai", "Comandante Anai"],
            blindados=["A'nai"], traduziveis=["Commander"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Comandante A\'nai"', '"Comandante A\'nai"'])
        self.assertEqual(result["reconstructed_names"], 1)
        self.assertEqual(result["protected_names"], 2)
        self.assertEqual(before, after)

    def test_overlap_em_nome_composto_nao_substitui_parte_de_outro(self):
        td, result, before, after, lines = self._pair(
            ["Commander Anna", "Commander Ann"],
            ["Comandante anna", "Comandante Ana"],
            blindados=["Ann", "Anna"], traduziveis=["Commander"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Comandante Anna"', '"Comandante Ann"'])
        self.assertEqual(result["reconstructed_names"], 1)
        self.assertEqual(result["protected_names"], 2)
        self.assertEqual(before, after)

    def test_nome_dentro_de_token_renpy_nao_e_texto_blindavel(self):
        td, result, before, after, lines = self._pair(
            ["[Gaius]"], ["[gaius]"],
            blindados=["Gaius"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"[gaius]"'])
        self.assertEqual(result["protected_names"], 0)
        self.assertEqual(result["unresolved_names"], [])
        self.assertEqual(before, after)

    def test_nome_em_dialogo_proximo_a_token_renpy_preserva_token(self):
        td, result, before, after, lines = self._pair(
            ["Commander [player] Gaius {b}"],
            ["Comandante [player] gaius {b}"],
            blindados=["Gaius"], traduziveis=["Commander"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Comandante [player] Gaius {b}"'])
        self.assertEqual(result["protected_names"], 1)
        self.assertEqual(result["unresolved_names"], [])
        self.assertEqual(before, after)

    def test_token_renpy_impede_reconstrucao_posicional_ambigua(self):
        td, result, before, after, lines = self._pair(
            ["Commander [player] Gaius"],
            ["Comandante [player] Go"],
            blindados=["Gaius"], traduziveis=["Commander"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Comandante [player] Go"'])
        self.assertEqual(result["reconstructed_names"], 0)
        self.assertEqual(len(result["unresolved_names"]), 1)
        self.assertEqual(before, after)

    def test_docx_reconstroi_nome_composto_preservando_nativo(self):
        td = tempfile.TemporaryDirectory()
        self.addCleanup(td.cleanup)
        root = Path(td.name)
        n, t, b = root / "nativo.docx", root / "traduzido.docx", root / "blindagem.txt"
        for path, value in ((n, '"Commander Gaius"'), (t, '"Comandante Go"')):
            doc = Document()
            doc.add_paragraph(value)
            doc.save(path)
        salvar_listas(b, {
            LISTA_BLINDADOS: ["Gaius"],
            LISTA_TRADUZIVEIS: ["Commander"],
        })
        before = n.read_bytes()
        result = review_json_pair(n, t, root, b)
        self.assertEqual([p.text for p in Document(t).paragraphs], ['"Comandante Gaius"'])
        self.assertEqual(result["reconstructed_names"], 1)
        self.assertEqual(result["unresolved_names"], [])
        self.assertEqual(n.read_bytes(), before)

    def test_docx_restaura_nome_integral_sem_alterar_nativo(self):
        td = tempfile.TemporaryDirectory()
        self.addCleanup(td.cleanup)
        root = Path(td.name)
        n = root / "nativo.docx"
        t = root / "traduzido.docx"
        b = root / "blindagem.txt"

        for path, value in ((n, '"Dalen"'), (t, '"O Vale"')):
            doc = Document()
            doc.add_paragraph(value)
            doc.save(path)

        salvar_listas(b, {
            LISTA_BLINDADOS: ["Dalen"],
            LISTA_TRADUZIVEIS: [],
        })
        before = n.read_bytes()
        result = review_json_pair(n, t, root, b)
        after = n.read_bytes()

        out = Document(t)
        lines = [p.text for p in out.paragraphs]
        self.assertEqual(lines, ['"Dalen"'])
        self.assertEqual(result["reconstructed_names"], 1)
        self.assertEqual(result["unresolved_names"], [])
        self.assertEqual(before, after)

    def test_conflito_nao_recebe_prioridade_automatica(self):
        td, result, before, after, lines = self._pair(
            ["Princess Christine"], ["Princesa Cristina"],
            blindados=["Christine"], traduziveis=["Christine"]
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Princesa Cristina"'])
        self.assertEqual(result["protected_names"], 0)
        self.assertEqual(result["conflict_records"], 1)
        self.assertEqual(before, after)


    def test_corpus_real_detecta_relacao_aberta_especificada(self):
        issues = diagnosticar_par_linguistico(
            "She's your younger... (Enter relationship - Default: [a_mc_rel])",
            "Ela é sua filha mais nova... (Insira o tipo de relacionamento - Padrão: [a_mc_rel])",
        )
        self.assertIn("relacao_aberta_especificada", {x["code"] for x in issues})

    def test_corpus_real_detecta_call_como_ligar(self):
        issues = diagnosticar_par_linguistico(
            "You call [a_name] and/or [b_name] your ...",
            "Você liga [a_name] e/ou [b_name] seu...",
        )
        self.assertIn("call_falso_cognato", {x["code"] for x in issues})

    def test_corpus_real_detecta_relacao_pronominal_quebrada(self):
        issues = diagnosticar_par_linguistico(
            "So I am her ...",
            "Então eu sou ela...",
        )
        self.assertIn("relacao_pronominal_quebrada", {x["code"] for x in issues})

    def test_review_pair_reporta_diagnostico_sem_retraduzir(self):
        td, result, before, after, lines = self._pair(
            ["She's your older..."],
            ["Ela é sua irmã mais velha..."],
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(lines, ['"Ela é sua irmã mais velha..."'])
        self.assertEqual(len(result["linguistic_issues"]), 1)
        self.assertEqual(result["linguistic_issues"][0]["code"], "relacao_aberta_especificada")
        self.assertEqual(before, after)

    def test_corpus_fisico_thief_of_hearts_variantes_reais(self):
        casos = (
            ("She's your younger... (Enter relationship - Default:", "Ela é sua mais nova... (Insira o tipo de relacionamento - Padrão:", set()),
            ("She's your younger...", "Ela é sua filha mais nova...", {"relacao_aberta_especificada"}),
            ("She's your older... (Enter relationship - Default:", "Ela é sua irmã mais velha... (Insira o tipo de relacionamento - Padrão:", {"relacao_aberta_especificada"}),
            ("She's your older...", "Ela é sua irmã mais velha...", {"relacao_aberta_especificada"}),
            ("So I am her ... (Enter relationship - Default:", "Então eu sou ela... (Inserir relação - Padrão:", {"relacao_pronominal_quebrada"}),
            ("So I am her ...", "Então eu sou ela...", {"relacao_pronominal_quebrada"}),
            ("You call [a_name] and/or [b_name] your ...", "Você liga [a_name] e/ou [b_name] seu...", {"call_falso_cognato"}),
            ("You call [v_name] and/or [vic_name] your ...", "Você liga [v_name] e/ou [vic_name] seu...", {"call_falso_cognato"}),
        )
        for native, translated, esperados in casos:
            with self.subTest(native=native, translated=translated):
                codigos = {x["code"] for x in diagnosticar_par_linguistico(native, translated)}
                self.assertEqual(codigos, esperados)

    def test_json_corrige_concordancia_atributo_sem_alterar_tags(self):
        native = "{b}{color=#87e1fa}Light Morality{/color}{/b} increased by {size=+3}{b}{color=#2bff8b}3{/color}{/b}{/size}."
        translated = "{b}{color=#87e1fa}Moralidade Leve{/color}{/b} aumentado por {size=+3}{b}{color=#2bff8b}3{/color}{/b}{/size}."
        actual, changes = _correct_source_anchored_structures(native, translated)
        self.assertIn("Moralidade Leve{/color}{/b} aumentada", actual)
        self.assertEqual(changes, 1)
        self.assertEqual(
            [s.valor for s in segmentar(native) if s.tipo == "protegido"],
            [s.valor for s in segmentar(actual) if s.tipo == "protegido"],
        )

    def test_json_reordena_placeholder_somente_com_possessivo_confirmado(self):
        native = "Looks like it's settled then, I agree with [mc]'s plan."
        translated = "Parece que está resolvido então, concordo com [mc] o plano de."
        actual, changes = _correct_source_anchored_structures(native, translated)
        self.assertEqual(actual, "Parece que está resolvido então, concordo com o plano de [mc].")
        self.assertEqual(changes, 1)

    def test_json_nao_reescreve_quando_original_nao_confirma_possessivo(self):
        native = "I agree with [mc] about the plan."
        translated = "concordo com [mc] o plano de."
        actual, changes = _correct_source_anchored_structures(native, translated)
        self.assertEqual(actual, translated)
        self.assertEqual(changes, 0)

if __name__ == "__main__":
    unittest.main()
